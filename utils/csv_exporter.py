import csv
import os
from datetime import datetime
from models.medicine_model import get_all_medicines


def export_inventory_csv(filepath=None):
    """
    Export all medicines to a CSV file.
    Returns the filepath on success, None on failure.
    """
    if not filepath:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filepath = f"medicine_inventory_{timestamp}.csv"

    medicines = get_all_medicines()
    if not medicines:
        print("[EXPORT] No medicines to export.")
        return None

    fieldnames = [
        'id', 'name', 'category_name', 'quantity', 'unit_name',
        'expiry_date', 'price', 'manufacturer_name', 'batch_number',
        'low_stock_threshold', 'created_at'
    ]

    try:
        with open(filepath, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction='ignore')
            writer.writeheader()
            writer.writerows(medicines)
        print(f"[EXPORT] Saved to {os.path.abspath(filepath)}")
        return filepath
    except Exception as e:
        print(f"[EXPORT ERROR] {e}")
        return None


def export_inventory_excel(filepath=None):
    """
    Export all medicines to an Excel file (.xlsx).
    Returns the filepath on success, None on failure.
    """
    if not filepath:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filepath = f"medicine_inventory_{timestamp}.xlsx"

    medicines = get_all_medicines()
    if not medicines:
        print("[EXPORT] No medicines to export.")
        return None

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory"

        headers = [
            'ID', 'Name', 'Category', 'Quantity', 'Unit',
            'Expiry Date', 'Price', 'Manufacturer', 'Batch Number',
            'Low Stock Threshold', 'Created At'
        ]
        keys = [
            'id', 'name', 'category_name', 'quantity', 'unit_name',
            'expiry_date', 'price', 'manufacturer_name', 'batch_number',
            'low_stock_threshold', 'created_at'
        ]

        # Header row styling
        header_fill = PatternFill("solid", fgColor="4F46E5")
        header_font = Font(bold=True, color="FFFFFF")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_idx, med in enumerate(medicines, 2):
            for col_idx, key in enumerate(keys, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(med.get(key, '')))

        # Auto column width
        for col in ws.columns:
            max_len = max((len(str(cell.value or '')) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        wb.save(filepath)
        print(f"[EXPORT] Saved to {os.path.abspath(filepath)}")
        return filepath
    except ImportError:
        print("[EXPORT] openpyxl not installed. Run: py -m pip install openpyxl")
        return None
    except Exception as e:
        print(f"[EXPORT ERROR] {e}")
        return None


def export_inventory_pdf(filepath=None):
    """
    Export all medicines to a PDF file.
    Returns the filepath on success, None on failure.
    """
    if not filepath:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filepath = f"medicine_inventory_{timestamp}.pdf"

    medicines = get_all_medicines()
    if not medicines:
        print("[EXPORT] No medicines to export.")
        return None

    try:
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        doc = SimpleDocTemplate(filepath, pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        elements = []

        # Title
        elements.append(Paragraph("Medicine Inventory Report", styles['Title']))
        elements.append(Paragraph(
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            styles['Normal']
        ))
        elements.append(Spacer(1, 12))

        # Table data
        headers = ['Name', 'Category', 'Qty', 'Unit', 'Expiry', 'Price', 'Manufacturer', 'Batch']
        keys    = ['name', 'category_name', 'quantity', 'unit_name', 'expiry_date', 'price', 'manufacturer_name', 'batch_number']

        data = [headers]
        for med in medicines:
            data.append([str(med.get(k, '')) for k in keys])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
            ('TEXTCOLOR',  (0, 0), (-1, 0), colors.white),
            ('FONTNAME',   (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0, 0), (-1, 0), 10),
            ('ALIGN',      (0, 0), (-1, -1), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('FONTSIZE',   (0, 1), (-1, -1), 9),
            ('GRID',       (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING',    (0, 0), (-1, -1), 6),
        ]))
        elements.append(table)

        doc.build(elements)
        print(f"[EXPORT] Saved to {os.path.abspath(filepath)}")
        return filepath
    except ImportError:
        print("[EXPORT] reportlab not installed. Run: py -m pip install reportlab")
        return None
    except Exception as e:
        print(f"[EXPORT ERROR] {e}")
        return None